from pathlib import Path
import re
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MD_PATH = ROOT / "2026_summer_internship_summary.md"
DOCX_PATH = ROOT / "2026_summer_internship_summary.docx"
NAVY, BLUE, LIGHT, GRAY = "000000", "000000", "EAF1F8", "666666"

def set_run_font(run, size=None, bold=None, color=None, east="Microsoft YaHei"):
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:eastAsia"), east)
    rpr.rFonts.set(qn("w:ascii"), "Calibri")
    rpr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_table_widths(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW")) or OxmlElement("w:tblW")
    if tbl_w.getparent() is None: tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths))); tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd")) or OxmlElement("w:tblInd")
    if tbl_ind.getparent() is None: tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120"); tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW")) or OxmlElement("w:tcW")
            if tc_w.getparent() is None: tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx])); tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

def add_inline(paragraph, text, size=10.5):
    for part in re.split(r"(\*\*.*?\*\*|`.*?`)", text):
        if not part: continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        clean = part[2:-2] if bold else (part[1:-1] if code else part)
        run = paragraph.add_run(clean)
        set_run_font(run, size, bold=bold)
        if code:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.15), Cm(2.0)
sec.left_margin = sec.right_margin = Cm(2.35)
sec.header_distance = sec.footer_distance = Cm(1.05)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"; normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(10.5); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.18
title_style = doc.styles["Title"]
title_style.font.name = "Calibri"; title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
title_style.font.size = Pt(28); title_style.font.bold = True; title_style.font.color.rgb = RGBColor(0, 0, 0)
for name, size, before, after, color in (
    ("Heading 1", 16, 16, 8, BLUE), ("Heading 2", 13, 12, 6, BLUE), ("Heading 3", 11.5, 8, 4, NAVY)):
    st = doc.styles[name]
    st.font.name = "Calibri"; st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

header = sec.header.paragraphs[0]
header.text = "中电福富 · 2026 暑期实习总结"
for run in header.runs: set_run_font(run, 8.5, True, NAVY)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run_font(footer.add_run("第 "), 9, color=GRAY)
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
set_run_font(footer.add_run(" 页"), 9, color=GRAY)

p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(96); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(p.add_run("INTERNSHIP REPORT"), 11, True, BLUE)
p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8)
set_run_font(p.add_run("2026 年暑期实习总结"), 28, True, NAVY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(30)
set_run_font(p.add_run("AI 应用开发 · Java Web 后端 · Hyperledger Fabric"), 14, color=BLUE)
for label, value in (("实习单位", "中电福富信息科技有限公司"), ("实习生", "杨企轩"), ("时间", "2026 年暑期")):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(5)
    set_run_font(p.add_run(f"{label}  |  "), 10.5, True, GRAY)
    set_run_font(p.add_run(value), 10.5, color=NAVY)
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(90); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run_font(p.add_run("基于项目源码、测试报告与验收记录整理"), 9.5, color=GRAY)
doc.add_page_break()

lines = MD_PATH.read_text(encoding="utf-8").splitlines()
i, title_skipped, body_started = 0, False, False
while i < len(lines):
    line = lines[i].rstrip()
    if line.startswith("# ") and not title_skipped:
        title_skipped = True; i += 1; continue
    if not body_started:
        if line.startswith("## "):
            body_started = True
        else:
            i += 1
            continue
    if not line: i += 1; continue
    if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|\s*[-: ]+", lines[i + 1]):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i += 1
        rows.pop(1)
        cols = len(rows[0]); table = doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = [2000, 3400, 3960] if cols == 3 else [int(9360 / cols)] * cols
        widths[-1] += 9360 - sum(widths); set_table_widths(table, widths)
        for rr, data in enumerate(rows):
            for cc, value in enumerate(data):
                cell = table.cell(rr, cc); cell.text = ""; cp = cell.paragraphs[0]
                cp.paragraph_format.space_after = Pt(2); cp.paragraph_format.line_spacing = 1.08
                add_inline(cp, value, 9.2)
                if rr == 0:
                    set_cell_shading(cell, LIGHT)
                    for run in cp.runs: run.bold = True; run.font.color.rgb = RGBColor.from_string(NAVY)
        doc.add_paragraph().paragraph_format.space_after = Pt(1)
        continue
    if line.startswith("### "):
        heading = re.sub(r"^\d+(?:\.\d+)*\s*", "", line[4:])
        doc.add_paragraph(heading, style="Heading 3")
    elif line.startswith("## "):
        heading = re.sub(r"^[\u4e00\u4e8c\u4e09\u56db\u4e94\u516d\u4e03\u516b\u4e5d\u5341]+\u3001", "", line[3:]).replace("\uff1a", " ")
        doc.add_paragraph(heading, style="Heading 1")
    elif re.match(r"^\d+\.\s+", line):
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent = Cm(.65)
        p.paragraph_format.first_line_indent = Cm(-.35); p.paragraph_format.space_after = Pt(4)
        add_inline(p, re.sub(r"^\d+\.\s+", "", line))
    elif line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.left_indent = Cm(.65)
        p.paragraph_format.first_line_indent = Cm(-.35); p.paragraph_format.space_after = Pt(4)
        add_inline(p, line[2:])
    else:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; add_inline(p, line)
    i += 1

doc.core_properties.title = "2026 年暑期实习总结"
doc.core_properties.subject = "中电福富实习：AI、Web 后端与 Hyperledger Fabric"
doc.core_properties.author = "杨企轩"
doc.save(DOCX_PATH)
print(DOCX_PATH)
